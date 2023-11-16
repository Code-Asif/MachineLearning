from docx import Document

def fill_rental_agreement_template(name, date, address, agreement_type):
    doc = Document('rentalAgreement.docx')

    placeholders = {
        'NamePlaceholder': name,
        'DatePlaceholder': date,
        'AddressPlaceholder': address,
        'AgreementTypePlaceholder': agreement_type,
        
    }

    for paragraph in doc.paragraphs:
        for placeholder, value in placeholders.items():
            if placeholder in paragraph.text:
                paragraph.text = paragraph.text.replace(placeholder, str(value))

    doc.save('NewRentalAgreement.docx')

if __name__ == "__main__":
    name = input("Enter your name: ")
    date = input("Enter the date: ")
    address = input("Enter your address: ")
    agreement_type = input("Enter the type of agreement: ")

    fill_rental_agreement_template(name, date, address, agreement_type)

    print("Filled rental agreement document saved as 'NewRentalAgreement.docx'")
